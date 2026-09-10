import os
import copy
import docx

def create_2025_qb_template():
    base_template_path = os.path.join(os.path.dirname(__file__), "..", "templates", "Question_Bank_Template.docx")
    out_path = os.path.join(os.path.dirname(__file__), "..", "templates", "Question_Bank_Template_2025.docx")
    
    doc = docx.Document(base_template_path)
    
    # 1. Fix typo in Unit V heading if present
    for p in doc.paragraphs:
        if p.text.strip().startswith("UNIT V:"):
            p.text = "UNIT V:"
            
    # 2. Update Regulation line to prefill 2025, leaving No. of Units blank
    for p in doc.paragraphs:
        if p.text.strip().startswith("Regulation"):
            p.text = "Regulation\t:\t2025\t\tNo. of Units\t:\t"

    # 3. Update existing section headings across Units I through V:
    # Under 2025 Regulation:
    # Part A (2 marks) -> PART A (1 mark)
    # Part B (13 marks) -> PART A (3 marks)
    # Part C (15 marks) -> PART B (10 marks)
    # There is NO Part C in 2025.
    for p in doc.paragraphs:
        txt = p.text.strip()
        if "PART A (2 marks)" in txt:
            p.text = "PART A (1 mark)"
        elif "PART B (13 marks)" in txt:
            p.text = "PART A (3 marks)"
        elif "PART C (15 marks)" in txt:
            p.text = "PART B (10 marks)"
            
    # 3b. Remove any stray 4-column orphaned tables (e.g. table 5 in Unit II)
    for tbl in list(doc.tables):
        if len(tbl.rows) > 0 and len(tbl.rows[0].cells) == 4 and tbl.rows[0].cells[0].text.strip().isdigit():
            tbl._element.getparent().remove(tbl._element)

    # 4. Locate insertion point before Instructions
    target_p = None
    for p in doc.paragraphs:
        if p.text.strip().startswith("Instructions:"):
            target_p = p
            break
            
    if not target_p:
        raise RuntimeError("Instructions paragraph not found in template.")
        
    # Deep copy Unit V tables (last 3 tables) for Unit VI
    t_a1 = copy.deepcopy(doc.tables[-3]._element)
    t_a3 = copy.deepcopy(doc.tables[-2]._element)
    t_b10 = copy.deepcopy(doc.tables[-1]._element)
    
    def insert_p_before(ref_p, text, style_name='Normal'):
        new_p = doc.add_paragraph(text, style=style_name)
        ref_p._p.addprevious(new_p._p)
        return new_p

    def insert_tbl_before(ref_p, tbl_element):
        ref_p._p.addprevious(tbl_element)

    # Insert Unit VI Section
    insert_p_before(target_p, "", "Normal")
    insert_p_before(target_p, "UNIT VI:", "Heading 1")
    insert_p_before(target_p, "", "Normal")

    # Unit VI Part A (1 mark)
    insert_p_before(target_p, "PART A (1 mark)", "Heading 2")
    insert_tbl_before(target_p, t_a1)
    insert_p_before(target_p, "", "Normal")

    # Unit VI Part A (3 marks)
    insert_p_before(target_p, "PART A (3 marks)", "Heading 2")
    insert_tbl_before(target_p, t_a3)
    insert_p_before(target_p, "", "Normal")

    # Unit VI Part B (10 marks)
    insert_p_before(target_p, "PART B (10 marks)", "Heading 2")
    insert_tbl_before(target_p, t_b10)
    insert_p_before(target_p, "", "Normal")

    # 5. Refine and update instruction points for 2025
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt.startswith("1. Header Information:"):
            p.text = "1. Header Information: Fill in the Department, Subject Code, Year/Sem, Subject Name, Regulation (2025), and No. of Units (4, 5, or 6) at the top of the question bank accurately."
        elif txt.startswith("2. Table Structure & Columns:"):
            p.text = "2. Table Structure & Columns: Each unit contains dedicated tables for PART A (1 Mark), PART A (3 Marks), and PART B (10 Marks). (Note: Under Regulation 2025, there is no Part C; the question paper comprises Part A (1 Mark), Part A (3 Marks), and Part B (10 Marks)). Enter Question S.No. in Column 1, Question Text in Column 2, and Knowledge Level in Column 3."
        elif txt.startswith("4. Course Outcomes"):
            p.text = "4. Course Outcomes (CO Mapping): Course Outcomes are automatically inherited from the parent Unit (Unit I \u2192 CO1, Unit II \u2192 CO2, Unit III \u2192 CO3, Unit IV \u2192 CO4, Unit V \u2192 CO5, Unit VI \u2192 CO6). You do not need to manually enter CO tags."

    # Add Instruction 7 explaining the 6-unit default and how to remove extra units
    instruction_7 = (
        "7. Unit Count Customization (4, 5, or 6 Units): This Question Bank template is pre-configured with 6 Units. "
        "If your course syllabus has only 4 or 5 Units, specify 'No. of Units: 4' (or 5) in the header, and simply delete "
        "the unused unit sections (e.g., Unit VI or Units V & VI) along with their respective tables before uploading. "
        "The system will automatically recognize the total unit count and apply the appropriate CAT examination unit-split rules."
    )
    doc.add_paragraph(instruction_7, style="Normal")

    doc.save(out_path)
    print(f"Regulation 2025 Question Bank Template generated successfully at: {out_path}")

if __name__ == "__main__":
    create_2025_qb_template()
