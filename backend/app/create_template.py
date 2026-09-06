import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set inner padding for table cells in dxa (1 pt = 20 dxa)."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m_name, m_val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m_name)
        node.set(qn('w:w'), str(m_val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex: str):
    """Set background color of a cell (e.g. 'F0F4F8')."""
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_table_borders(table, color="CCCCCC", sz="4", val="single"):
    """Apply clean subtle borders to the table."""
    tblPr = table._tbl.tblPr
    borders_xml = f'''
    <w:tblBorders {nsdecls("w")}>
        <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
    </w:tblBorders>
    '''
    tblPr.append(parse_xml(borders_xml))

def build_question_bank_template(output_path: str):
    doc = docx.Document()

    # --- PAGE MARGINS ---
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        section.page_width = Inches(8.27)   # A4 width
        section.page_height = Inches(11.69) # A4 height

    # Base Colors
    NAVY = RGBColor(0x1B, 0x36, 0x5D)      # Institutional Deep Navy
    DARK_GRAY = RGBColor(0x33, 0x33, 0x33) # Body Text
    MUTED_GRAY = RGBColor(0x66, 0x66, 0x66)
    HEADER_BG = "EBF2FA"                   # Light Soft Blue header background
    META_BG = "F8FAFC"                     # Light Gray/Slate
    BORDER_COLOR = "CBD5E1"                # Slate border

    # --- 1. INSTITUTION & DOCUMENT HEADER ---
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(2)
    r_dept = title_p.add_run("DEPARTMENT OF __________________________________\n")
    r_dept.font.name = "Calibri"
    r_dept.font.size = Pt(13)
    r_dept.font.bold = True
    r_dept.font.color.rgb = NAVY

    r_main = title_p.add_run("QUESTION BANK (ACADEMIC YEAR: 20___ - 20___) [ODD / EVEN] SEMESTER")
    r_main.font.name = "Calibri"
    r_main.font.size = Pt(12)
    r_main.font.bold = True
    r_main.font.color.rgb = DARK_GRAY

    # --- 2. METADATA TABLE (Blank details for faculty to fill in) ---
    meta_table = doc.add_table(rows=3, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(meta_table, color=BORDER_COLOR, sz="6")
    
    meta_data = [
        ("Department: ________________________", "Subject Code: ________________________"),
        ("Year / Semester: ____________________", "Subject Name: ________________________"),
        ("Regulation: 2021 / 2025              ", "Staff In-charge: ______________________"),
    ]

    for r_idx, (col1_text, col2_text) in enumerate(meta_data):
        row = meta_table.rows[r_idx]
        
        # Col 1
        c1 = row.cells[0]
        c1.width = Inches(3.38)
        set_cell_margins(c1, top=80, bottom=80, left=120, right=120)
        set_cell_shading(c1, META_BG)
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(0)
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(col1_text)
        r1.font.name = "Calibri"
        r1.font.size = Pt(10)
        r1.font.bold = True
        r1.font.color.rgb = DARK_GRAY

        # Col 2
        c2 = row.cells[1]
        c2.width = Inches(3.38)
        set_cell_margins(c2, top=80, bottom=80, left=120, right=120)
        set_cell_shading(c2, META_BG)
        p2 = c2.paragraphs[0]
        p2.paragraph_format.space_before = Pt(0)
        p2.paragraph_format.space_after = Pt(0)
        r2 = p2.add_run(col2_text)
        r2.font.name = "Calibri"
        r2.font.size = Pt(10)
        r2.font.bold = True
        r2.font.color.rgb = DARK_GRAY

    # Spacing
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # --- 3. INSTRUCTION / GUIDELINES BOX FOR FACULTY ---
    guide_table = doc.add_table(rows=1, cols=1)
    guide_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(guide_table, color="93C5FD", sz="6")
    g_cell = guide_table.rows[0].cells[0]
    g_cell.width = Inches(6.77)
    set_cell_margins(g_cell, top=100, bottom=100, left=140, right=140)
    set_cell_shading(g_cell, "F0F7FF")
    
    gp = g_cell.paragraphs[0]
    gp.paragraph_format.space_before = Pt(0)
    gp.paragraph_format.space_after = Pt(2)
    gr_head = gp.add_run("Faculty Instructions & Guidelines for Question Bank Entry:\n")
    gr_head.font.name = "Calibri"
    gr_head.font.size = Pt(9.5)
    gr_head.font.bold = True
    gr_head.font.color.rgb = NAVY

    guidelines_text = (
        "1. Fill in the Department, Subject Code, Subject Name, Semester, Regulation, and Staff In-charge in the table above.\n"
        "2. Knowledge Level (BT Level): Enter K1 (Remember), K2 (Understand), K3 (Apply), K4 (Analyze), K5 (Evaluate), K6 (Create).\n"
        "3. Course Outcome (CO): Enter CO1 (Unit 1), CO2 (Unit 2), CO3 (Unit 3), CO4 (Unit 4), CO5 (Unit 5).\n"
        "4. You can add or remove rows in each table as needed. Maintain the 4-column format: S. No | Question | Knowledge Level | Course Outcome."
    )
    gr_body = gp.add_run(guidelines_text)
    gr_body.font.name = "Calibri"
    gr_body.font.size = Pt(8.5)
    gr_body.font.color.rgb = DARK_GRAY

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Helper function to create Question tables for parts
    def create_part_table(part_title: str, marks_label: str, num_sample_rows: int, default_co: str, default_kl: str):
        # Section Heading
        head_p = doc.add_paragraph()
        head_p.paragraph_format.space_before = Pt(8)
        head_p.paragraph_format.space_after = Pt(3)
        r_part = head_p.add_run(f"{part_title} ({marks_label})")
        r_part.font.name = "Calibri"
        r_part.font.size = Pt(11)
        r_part.font.bold = True
        r_part.font.color.rgb = NAVY

        # 4-Column Table
        table = doc.add_table(rows=num_sample_rows + 1, cols=4)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(table, color=BORDER_COLOR, sz="4")

        col_widths = [Inches(0.6), Inches(4.37), Inches(0.9), Inches(0.9)]
        headers = ["S. No", "Question", "Knowledge Level", "Course Outcome"]

        # Format Header Row
        hdr_row = table.rows[0]
        trPr = hdr_row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

        for c_idx, title in enumerate(headers):
            cell = hdr_row.cells[c_idx]
            cell.width = col_widths[c_idx]
            set_cell_margins(cell, top=80, bottom=80, left=60, right=60)
            set_cell_shading(cell, HEADER_BG)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx != 1 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(title)
            r.font.name = "Calibri"
            r.font.size = Pt(9.5)
            r.font.bold = True
            r.font.color.rgb = NAVY

        # Populate Blank / Placeholder Rows
        for r_idx in range(1, num_sample_rows + 1):
            row = table.rows[r_idx]
            # S.No
            c0 = row.cells[0]
            c0.width = col_widths[0]
            set_cell_margins(c0, top=60, bottom=60, left=40, right=40)
            c0.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p0 = c0.paragraphs[0]
            p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p0.paragraph_format.space_before = Pt(0)
            p0.paragraph_format.space_after = Pt(0)
            r0 = p0.add_run(str(r_idx))
            r0.font.name = "Calibri"
            r0.font.size = Pt(9.5)

            # Question Text Cell (Blank / Click to type)
            c1 = row.cells[1]
            c1.width = col_widths[1]
            set_cell_margins(c1, top=60, bottom=60, left=60, right=60)
            c1.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p1 = c1.paragraphs[0]
            p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p1.paragraph_format.space_before = Pt(0)
            p1.paragraph_format.space_after = Pt(0)
            if r_idx == 1:
                r1 = p1.add_run("[Enter question text here...]")
                r1.font.italic = True
                r1.font.color.rgb = MUTED_GRAY
            else:
                r1 = p1.add_run("")
            r1.font.name = "Calibri"
            r1.font.size = Pt(9.5)

            # Knowledge Level
            c2 = row.cells[2]
            c2.width = col_widths[2]
            set_cell_margins(c2, top=60, bottom=60, left=40, right=40)
            c2.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p2 = c2.paragraphs[0]
            p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p2.paragraph_format.space_before = Pt(0)
            p2.paragraph_format.space_after = Pt(0)
            r2 = p2.add_run(default_kl)
            r2.font.name = "Calibri"
            r2.font.size = Pt(9.5)
            r2.font.color.rgb = DARK_GRAY

            # Course Outcome
            c3 = row.cells[3]
            c3.width = col_widths[3]
            set_cell_margins(c3, top=60, bottom=60, left=40, right=40)
            c3.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p3 = c3.paragraphs[0]
            p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p3.paragraph_format.space_before = Pt(0)
            p3.paragraph_format.space_after = Pt(0)
            r3 = p3.add_run(default_co)
            r3.font.name = "Calibri"
            r3.font.size = Pt(9.5)
            r3.font.color.rgb = DARK_GRAY

        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # --- 4. UNITS (UNIT I TO UNIT V) ---
    units = [
        ("UNIT I", "CO1", "INTRODUCTION / UNIT 1 TITLE"),
        ("UNIT II", "CO2", "UNIT 2 TITLE"),
        ("UNIT III", "CO3", "UNIT 3 TITLE"),
        ("UNIT IV", "CO4", "UNIT 4 TITLE"),
        ("UNIT V", "CO5", "UNIT 5 TITLE"),
    ]

    for u_idx, (unit_label, co_label, unit_name) in enumerate(units):
        # Unit Header Box
        unit_p = doc.add_paragraph()
        unit_p.paragraph_format.space_before = Pt(12)
        unit_p.paragraph_format.space_after = Pt(2)
        r_u_num = unit_p.add_run(f"{unit_label}: ")
        r_u_num.font.name = "Calibri"
        r_u_num.font.size = Pt(12)
        r_u_num.font.bold = True
        r_u_num.font.color.rgb = NAVY

        r_u_name = unit_p.add_run(f"_____________________________________________")
        r_u_name.font.name = "Calibri"
        r_u_name.font.size = Pt(12)
        r_u_name.font.bold = True
        r_u_name.font.color.rgb = DARK_GRAY

        # Syllabus line
        syl_p = doc.add_paragraph()
        syl_p.paragraph_format.space_before = Pt(0)
        syl_p.paragraph_format.space_after = Pt(6)
        r_syl = syl_p.add_run("[Enter syllabus topics / description here for " + unit_label + "]")
        r_syl.font.name = "Calibri"
        r_syl.font.size = Pt(9)
        r_syl.font.italic = True
        r_syl.font.color.rgb = MUTED_GRAY

        # Part A (2 Marks / 1 Mark)
        create_part_table(
            part_title="Part A",
            marks_label="Two Marks Questions",
            num_sample_rows=5,
            default_co=co_label,
            default_kl="K1"
        )

        # Part B (13 Marks / 16 Marks)
        create_part_table(
            part_title="Part B",
            marks_label="Thirteen / Sixteen Marks Questions",
            num_sample_rows=3,
            default_co=co_label,
            default_kl="K2"
        )

        # Part C (15 Marks / 12 Marks)
        create_part_table(
            part_title="Part C",
            marks_label="Fifteen Marks Questions (Application / Design / Case Study)",
            num_sample_rows=2,
            default_co=co_label,
            default_kl="K3"
        )

        # Page break between units for clean print formatting
        if u_idx < len(units) - 1:
            doc.add_page_break()

    # --- 5. SIGNATURE BLOCK AT END ---
    doc.add_paragraph().paragraph_format.space_before = Pt(20)
    sig_p = doc.add_paragraph()
    sig_p.paragraph_format.space_before = Pt(20)
    sig_p.paragraph_format.space_after = Pt(0)
    
    r_sig = sig_p.add_run(
        "Staff In-Charge: ______________________                     HOD / Principal: ______________________"
    )
    r_sig.font.name = "Calibri"
    r_sig.font.size = Pt(11)
    r_sig.font.bold = True
    r_sig.font.color.rgb = DARK_GRAY

    # Save Document
    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    doc.save(output_path)
    print(f"Successfully generated Question Bank Template at: {output_path}")

if __name__ == "__main__":
    out_file = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "templates", "Question_Bank_Template.docx"))
    build_question_bank_template(out_file)
