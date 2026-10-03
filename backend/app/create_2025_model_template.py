import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_cell_border(cell, **kwargs):
    """
    kwargs: top, bottom, left, right
    values: dict(val='single', sz='4', color='000000', space='0')
    """
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = tcPr.find(qn('w:tcBorders'))
    if tcBorders is None:
        tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}/>')
        tcPr.append(tcBorders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge_data = kwargs.get(edge)
        if edge_data:
            val = edge_data.get('val', 'single')
            sz = edge_data.get('sz', '4')
            color = edge_data.get('color', '000000')
            space = edge_data.get('space', '0')
            b_elem = parse_xml(f'<w:{edge} {nsdecls("w")} w:val="{val}" w:sz="{sz}" w:space="{space}" w:color="{color}"/>')
            tcBorders.append(b_elem)

def set_table_borders(table, color="000000", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    tblBorders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        </w:tblBorders>
    ''')
    tblPr.append(tblBorders)

def create_2025_model_template(output_path: str):
    doc = docx.Document()
    
    # 1. Section Margins (A4)
    sec = doc.sections[0]
    sec.page_width = Inches(8.27)
    sec.page_height = Inches(11.69)
    sec.top_margin = Inches(0.75)
    sec.bottom_margin = Inches(0.5)
    sec.left_margin = Inches(0.75)
    sec.right_margin = Inches(0.65)
    
    # 2. Top SET Indicator
    p_set = doc.add_paragraph()
    p_set.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_set.paragraph_format.space_before = Pt(0)
    p_set.paragraph_format.space_after = Pt(2)
    r_set = p_set.add_run("SET - I")
    r_set.font.name = "Times New Roman"
    r_set.font.size = Pt(11)
    r_set.bold = True
    
    # 3. Institution Header Table (Table 0)
    t0 = doc.add_table(rows=2, cols=1)
    t0.alignment = WD_TABLE_ALIGNMENT.CENTER
    t0.autofit = False
    set_table_borders(t0, color="000000", sz="6", val="single")
    
    c00 = t0.rows[0].cells[0]
    c00.width = Inches(6.87)
    p00 = c00.paragraphs[0]
    p00.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p00.paragraph_format.space_before = Pt(2)
    p00.paragraph_format.space_after = Pt(1)
    r_inst = p00.add_run("NAME OF THE INSTITUTION :__________________________________________")
    r_inst.font.name = "Times New Roman"
    r_inst.font.size = Pt(12)
    r_inst.bold = True
    
    c01 = t0.rows[1].cells[0]
    c01.width = Inches(6.87)
    p01 = c01.paragraphs[0]
    p01.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p01.paragraph_format.space_before = Pt(1)
    p01.paragraph_format.space_after = Pt(3)
    r_subinst = p01.add_run("(Approved by AICTE, Affiliated to Anna University Chennai & NAAC Accredited Institution)\nChennai, Tamil Nadu.")
    r_subinst.font.name = "Times New Roman"
    r_subinst.font.size = Pt(10)
    
    # 4. Date Paragraph
    p_date = doc.add_paragraph()
    p_date.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_date.paragraph_format.space_before = Pt(4)
    p_date.paragraph_format.space_after = Pt(2)
    r_date = p_date.add_run("Date: ______________")
    r_date.font.name = "Times New Roman"
    r_date.font.size = Pt(11)
    r_date.bold = True
    
    # 5. Exam Title & Metadata Paragraphs
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(2)
    p_title.paragraph_format.space_after = Pt(1)
    r_title = p_title.add_run("MODEL EXAMINATION")
    r_title.font.name = "Times New Roman"
    r_title.font.size = Pt(14)
    r_title.bold = True
    r_title.underline = True
    
    p_reg = doc.add_paragraph()
    p_reg.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_reg.paragraph_format.space_before = Pt(1)
    p_reg.paragraph_format.space_after = Pt(1)
    r_reg = p_reg.add_run("(2025-Regulation)")
    r_reg.font.name = "Times New Roman"
    r_reg.font.size = Pt(12)
    r_reg.bold = True
    
    p_sem = doc.add_paragraph()
    p_sem.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sem.paragraph_format.space_before = Pt(1)
    p_sem.paragraph_format.space_after = Pt(4)
    r_sem = p_sem.add_run("ODD SEMESTER-2025-26")
    r_sem.font.name = "Times New Roman"
    r_sem.font.size = Pt(12)
    
    # 6. Course Details Paragraphs
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(2)
    p_sub.paragraph_format.space_after = Pt(2)
    r_sub1 = p_sub.add_run("Sub. Code/Sub.Name: ")
    r_sub1.font.name = "Times New Roman"
    r_sub1.font.size = Pt(11)
    r_sub1.bold = True
    r_sub2 = p_sub.add_run("OCS353/ Data Science fundamentals")
    r_sub2.font.name = "Times New Roman"
    r_sub2.font.size = Pt(11)
    
    p_deg = doc.add_paragraph()
    p_deg.paragraph_format.space_before = Pt(2)
    p_deg.paragraph_format.space_after = Pt(2)
    r_deg1 = p_deg.add_run("Degree/Branch/Sem: ")
    r_deg1.font.name = "Times New Roman"
    r_deg1.font.size = Pt(11)
    r_deg1.bold = True
    r_deg2 = p_deg.add_run("BE/BTECH/ CIVIL/AERO/MECH/EEE/TEXT/VII")
    r_deg2.font.name = "Times New Roman"
    r_deg2.font.size = Pt(11)
    
    p_time = doc.add_paragraph()
    p_time.paragraph_format.space_before = Pt(2)
    p_time.paragraph_format.space_after = Pt(4)
    r_t1 = p_time.add_run("Time : 3 Hours\t\t\t\t\tMaximum Marks: 100")
    r_t1.font.name = "Times New Roman"
    r_t1.font.size = Pt(11)
    r_t1.bold = True
    
    # 7. Knowledge Level Table (Table 1)
    t_kl = doc.add_table(rows=1, cols=1)
    t_kl.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_kl.autofit = False
    set_table_borders(t_kl, color="666666", sz="4", val="single")
    c_kl = t_kl.rows[0].cells[0]
    c_kl.width = Inches(6.87)
    set_cell_margins(c_kl, top=60, bottom=60, left=80, right=80)
    p_kl = c_kl.paragraphs[0]
    p_kl.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_kl.paragraph_format.space_before = Pt(1)
    p_kl.paragraph_format.space_after = Pt(1)
    r_kl = p_kl.add_run(
        "Knowledge Level: K1–Remember (Define, List, State, Identify, Recall, Name, Mention)   "
        "K2–Understand (Explain, Describe, Discuss, Distinguish, Illustrate)   "
        "K3–Apply (Compute, Calculate, Solve, Apply, Derive, Demonstrate, Determine)   "
        "K4–Analyze (Analyze, Differentiate, Examine, Classify, Compare, Investigate)   "
        "K5–Evaluate (Justify, Evaluate, Assess, Critique, Validate)   "
        "K6–Create (Design, Develop, Construct, Formulate, Propose)"
    )
    r_kl.font.name = "Times New Roman"
    r_kl.font.size = Pt(8.5)
    r_kl.font.italic = True
    
    # Spacing
    p_sp1 = doc.add_paragraph()
    p_sp1.paragraph_format.space_before = Pt(2)
    p_sp1.paragraph_format.space_after = Pt(2)
    
    # 8. PART A (Section 1: 10 x 1 = 10 Marks - MCQ)
    p_pa1_title = doc.add_paragraph()
    p_pa1_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_pa1_title.paragraph_format.space_before = Pt(4)
    p_pa1_title.paragraph_format.space_after = Pt(1)
    r_pa1_t = p_pa1_title.add_run("PART – A (10 x 1 = 10 Marks)")
    r_pa1_t.font.name = "Times New Roman"
    r_pa1_t.font.size = Pt(13)
    r_pa1_t.bold = True
    
    p_pa1_sub = doc.add_paragraph()
    p_pa1_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_pa1_sub.paragraph_format.space_before = Pt(0)
    p_pa1_sub.paragraph_format.space_after = Pt(4)
    r_pa1_s = p_pa1_sub.add_run("Answer ALL the questions (Multiple Choice Questions)")
    r_pa1_s.font.name = "Times New Roman"
    r_pa1_s.font.size = Pt(10.5)
    r_pa1_s.italic = True
    
    # Table 2: Part A (1 Mark MCQ) [11 rows, 4 cols]
    col_widths_4 = [Inches(0.55), Inches(4.55), Inches(0.72), Inches(1.05)]
    t_pa1 = doc.add_table(rows=11, cols=4)
    t_pa1.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_pa1.autofit = False
    set_table_borders(t_pa1, color="000000", sz="4", val="single")
    
    # Header row
    hdr_titles_4 = ["Q.No.", "Questions", "KL", "CO Attainment"]
    for c_idx, title in enumerate(hdr_titles_4):
        cell = t_pa1.rows[0].cells[c_idx]
        cell.width = col_widths_4[c_idx]
        set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(title)
        r.font.name = "Times New Roman"
        r.font.size = Pt(11)
        r.bold = True
        
    for q_idx in range(10):
        row = t_pa1.rows[1 + q_idx]
        for c_idx in range(4):
            row.cells[c_idx].width = col_widths_4[c_idx]
            set_cell_margins(row.cells[c_idx], top=80, bottom=80, left=100, right=100)
        
        # Q.No
        p0 = row.cells[0].paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r0 = p0.add_run(f"{q_idx + 1}.")
        r0.font.name = "Times New Roman"
        r0.font.size = Pt(11)
        
        # KL
        p2 = row.cells[2].paragraphs[0]
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r2 = p2.add_run("K1" if q_idx % 2 == 0 else "K2")
        r2.font.name = "Times New Roman"
        r2.font.size = Pt(11)
        
        # CO
        p3 = row.cells[3].paragraphs[0]
        p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r3 = p3.add_run(f"CO{(q_idx // 2) + 1}")
        r3.font.name = "Times New Roman"
        r3.font.size = Pt(11)

    # Spacing
    p_sp2 = doc.add_paragraph()
    p_sp2.paragraph_format.space_before = Pt(4)
    p_sp2.paragraph_format.space_after = Pt(2)

    # 9. PART A (Section 2: 10 x 3 = 30 Marks - Short Answer)
    p_pa3_title = doc.add_paragraph()
    p_pa3_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_pa3_title.paragraph_format.space_before = Pt(4)
    p_pa3_title.paragraph_format.space_after = Pt(1)
    r_pa3_t = p_pa3_title.add_run("PART – A (10 x 3 = 30 Marks)")
    r_pa3_t.font.name = "Times New Roman"
    r_pa3_t.font.size = Pt(13)
    r_pa3_t.bold = True
    
    p_pa3_sub = doc.add_paragraph()
    p_pa3_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_pa3_sub.paragraph_format.space_before = Pt(0)
    p_pa3_sub.paragraph_format.space_after = Pt(4)
    r_pa3_s = p_pa3_sub.add_run("Answer ALL the questions (Short Answer Questions)")
    r_pa3_s.font.name = "Times New Roman"
    r_pa3_s.font.size = Pt(10.5)
    r_pa3_s.italic = True

    # Table 3: Part A (3 Marks Short Answer) [11 rows, 4 cols]
    t_pa3 = doc.add_table(rows=11, cols=4)
    t_pa3.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_pa3.autofit = False
    set_table_borders(t_pa3, color="000000", sz="4", val="single")
    
    for c_idx, title in enumerate(hdr_titles_4):
        cell = t_pa3.rows[0].cells[c_idx]
        cell.width = col_widths_4[c_idx]
        set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(title)
        r.font.name = "Times New Roman"
        r.font.size = Pt(11)
        r.bold = True
        
    for q_idx in range(10):
        row = t_pa3.rows[1 + q_idx]
        for c_idx in range(4):
            row.cells[c_idx].width = col_widths_4[c_idx]
            set_cell_margins(row.cells[c_idx], top=80, bottom=80, left=100, right=100)
        
        # Q.No (11 to 20)
        p0 = row.cells[0].paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r0 = p0.add_run(f"{11 + q_idx}.")
        r0.font.name = "Times New Roman"
        r0.font.size = Pt(11)
        
        # KL
        p2 = row.cells[2].paragraphs[0]
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r2 = p2.add_run("K2" if q_idx % 2 == 0 else "K3")
        r2.font.name = "Times New Roman"
        r2.font.size = Pt(11)
        
        # CO
        p3 = row.cells[3].paragraphs[0]
        p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r3 = p3.add_run(f"CO{(q_idx // 2) + 1}")
        r3.font.name = "Times New Roman"
        r3.font.size = Pt(11)

    # 10. PART B (5 x 12 = 60 Marks - Either OR Type)
    p_pb_title = doc.add_paragraph()
    p_pb_title.paragraph_format.page_break_before = True
    p_pb_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_pb_title.paragraph_format.space_before = Pt(6)
    p_pb_title.paragraph_format.space_after = Pt(1)
    r_pb_t = p_pb_title.add_run("PART – B (5 x 12 = 60 Marks)")
    r_pb_t.font.name = "Times New Roman"
    r_pb_t.font.size = Pt(13)
    r_pb_t.bold = True
    
    p_pb_sub = doc.add_paragraph()
    p_pb_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_pb_sub.paragraph_format.space_before = Pt(0)
    p_pb_sub.paragraph_format.space_after = Pt(4)
    r_pb_s = p_pb_sub.add_run("Answer ALL the questions (Either OR Type)")
    r_pb_s.font.name = "Times New Roman"
    r_pb_s.font.size = Pt(10.5)
    r_pb_s.italic = True

    # Table 4: Part B (5 Either-Or Pairs, 16 rows: Header + 5 * 3 rows, 5 cols)
    col_widths_5 = [Inches(0.50), Inches(0.40), Inches(4.20), Inches(0.72), Inches(1.05)]
    t_pb = doc.add_table(rows=16, cols=5)
    t_pb.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_pb.autofit = False
    set_table_borders(t_pb, color="000000", sz="4", val="single")
    
    hdr_titles_5 = ["Q.No.", "Sub", "Questions", "KL", "CO Attainment"]
    for c_idx, title in enumerate(hdr_titles_5):
        cell = t_pb.rows[0].cells[c_idx]
        cell.width = col_widths_5[c_idx]
        set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(title)
        r.font.name = "Times New Roman"
        r.font.size = Pt(11)
        r.bold = True

    for q_idx in range(5):
        q_num = 21 + q_idx
        row_a_idx = 1 + q_idx * 3
        row_or_idx = 2 + q_idx * 3
        row_b_idx = 3 + q_idx * 3
        
        # Setup row A
        row_a = t_pb.rows[row_a_idx]
        for c_idx in range(5):
            row_a.cells[c_idx].width = col_widths_5[c_idx]
            set_cell_margins(row_a.cells[c_idx], top=80, bottom=80, left=100, right=100)
        p0 = row_a.cells[0].paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r0 = p0.add_run(f"{q_num}.")
        r0.font.name = "Times New Roman"
        r0.font.size = Pt(11)
        r0.bold = True
        
        p1 = row_a.cells[1].paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r1 = p1.add_run("(a)")
        r1.font.name = "Times New Roman"
        r1.font.size = Pt(11)
        
        p3 = row_a.cells[3].paragraphs[0]
        p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r3 = p3.add_run("K3")
        r3.font.name = "Times New Roman"
        r3.font.size = Pt(11)
        
        p4 = row_a.cells[4].paragraphs[0]
        p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r4 = p4.add_run(f"CO{q_idx + 1}")
        r4.font.name = "Times New Roman"
        r4.font.size = Pt(11)

        # Setup row OR
        row_or = t_pb.rows[row_or_idx]
        for c_idx in range(5):
            row_or.cells[c_idx].width = col_widths_5[c_idx]
            set_cell_margins(row_or.cells[c_idx], top=40, bottom=40, left=100, right=100)
        p_or = row_or.cells[2].paragraphs[0]
        p_or.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_or = p_or.add_run("(OR)")
        r_or.font.name = "Times New Roman"
        r_or.font.size = Pt(10.5)
        r_or.bold = True

        # Setup row B
        row_b = t_pb.rows[row_b_idx]
        for c_idx in range(5):
            row_b.cells[c_idx].width = col_widths_5[c_idx]
            set_cell_margins(row_b.cells[c_idx], top=80, bottom=80, left=100, right=100)
        p1_b = row_b.cells[1].paragraphs[0]
        p1_b.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r1_b = p1_b.add_run("(b)")
        r1_b.font.name = "Times New Roman"
        r1_b.font.size = Pt(11)
        
        p3_b = row_b.cells[3].paragraphs[0]
        p3_b.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r3_b = p3_b.add_run("K3")
        r3_b.font.name = "Times New Roman"
        r3_b.font.size = Pt(11)
        
        p4_b = row_b.cells[4].paragraphs[0]
        p4_b.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r4_b = p4_b.add_run(f"CO{q_idx + 1}")
        r4_b.font.name = "Times New Roman"
        r4_b.font.size = Pt(11)

    # 11. Table of Specifications (Question – Wise) (Table 5)
    p_tos_q = doc.add_paragraph()
    p_tos_q.paragraph_format.page_break_before = True
    p_tos_q.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_tos_q.paragraph_format.space_before = Pt(6)
    p_tos_q.paragraph_format.space_after = Pt(4)
    r_tos_q = p_tos_q.add_run("Table of Specifications (Question – Wise)")
    r_tos_q.font.name = "Times New Roman"
    r_tos_q.font.size = Pt(12)
    r_tos_q.bold = True

    t_tos_q = doc.add_table(rows=8, cols=8)
    t_tos_q.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_tos_q.autofit = False
    set_table_borders(t_tos_q, color="000000", sz="4", val="single")
    
    col_widths_tos = [Inches(1.00), Inches(0.84), Inches(0.84), Inches(0.84), Inches(0.84), Inches(0.84), Inches(0.84), Inches(0.83)]
    
    # Row 0: SYLLABUS, No. Of Questions (spans across), Total
    row0_q = t_tos_q.rows[0]
    for c_idx, cell in enumerate(row0_q.cells):
        cell.width = col_widths_tos[c_idx]
        set_cell_margins(cell, top=60, bottom=60, left=60, right=60)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if c_idx == 0:
            r = p.add_run("SYLLABUS")
        elif c_idx == 7:
            r = p.add_run("Total")
        else:
            r = p.add_run("No. Of Questions")
        r.font.name = "Times New Roman"
        r.font.size = Pt(10)
        r.bold = True

    # Row 1: Unit, K1 Remembering, K2 Understanding, K3 Applying, K4 Analyzing, K5 Evaluating, K6 Creating, Total
    kl_labels = ["Unit", "K1 Remembering", "K2 Understanding", "K3 Applying", "K4 Analyzing", "K5 Evaluating", "K6 Creating", "Total"]
    row1_q = t_tos_q.rows[1]
    for c_idx, cell in enumerate(row1_q.cells):
        cell.width = col_widths_tos[c_idx]
        set_cell_margins(cell, top=60, bottom=60, left=60, right=60)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(kl_labels[c_idx])
        r.font.name = "Times New Roman"
        r.font.size = Pt(9.5)
        r.bold = True

    unit_names = ["Unit I", "Unit II", "Unit III", "Unit IV", "Unit V"]
    for u_idx, uname in enumerate(unit_names):
        row = t_tos_q.rows[2 + u_idx]
        for c_idx, cell in enumerate(row.cells):
            cell.width = col_widths_tos[c_idx]
            set_cell_margins(cell, top=60, bottom=60, left=60, right=60)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if c_idx == 0:
                r = p.add_run(uname)
                r.bold = True
            else:
                r = p.add_run("")
            r.font.name = "Times New Roman"
            r.font.size = Pt(10)

    # Total Row
    row_tot_q = t_tos_q.rows[7]
    for c_idx, cell in enumerate(row_tot_q.cells):
        cell.width = col_widths_tos[c_idx]
        set_cell_margins(cell, top=60, bottom=60, left=60, right=60)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if c_idx == 0:
            r = p.add_run("Total")
            r.bold = True
        else:
            r = p.add_run("")
            r.bold = True
        r.font.name = "Times New Roman"
        r.font.size = Pt(10)

    # 12. Table of Specifications (Marks – Wise) (Table 6)
    p_tos_m = doc.add_paragraph()
    p_tos_m.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_tos_m.paragraph_format.space_before = Pt(8)
    p_tos_m.paragraph_format.space_after = Pt(4)
    r_tos_m = p_tos_m.add_run("Table of Specifications (Marks – Wise)")
    r_tos_m.font.name = "Times New Roman"
    r_tos_m.font.size = Pt(12)
    r_tos_m.bold = True

    t_tos_m = doc.add_table(rows=8, cols=8)
    t_tos_m.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_tos_m.autofit = False
    set_table_borders(t_tos_m, color="000000", sz="4", val="single")

    row0_m = t_tos_m.rows[0]
    for c_idx, cell in enumerate(row0_m.cells):
        cell.width = col_widths_tos[c_idx]
        set_cell_margins(cell, top=60, bottom=60, left=60, right=60)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if c_idx == 0:
            r = p.add_run("SYLLABUS")
        elif c_idx == 7:
            r = p.add_run("Total")
        else:
            r = p.add_run("Marks")
        r.font.name = "Times New Roman"
        r.font.size = Pt(10)
        r.bold = True

    row1_m = t_tos_m.rows[1]
    for c_idx, cell in enumerate(row1_m.cells):
        cell.width = col_widths_tos[c_idx]
        set_cell_margins(cell, top=60, bottom=60, left=60, right=60)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(kl_labels[c_idx])
        r.font.name = "Times New Roman"
        r.font.size = Pt(9.5)
        r.bold = True

    for u_idx, uname in enumerate(unit_names):
        row = t_tos_m.rows[2 + u_idx]
        for c_idx, cell in enumerate(row.cells):
            cell.width = col_widths_tos[c_idx]
            set_cell_margins(cell, top=60, bottom=60, left=60, right=60)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if c_idx == 0:
                r = p.add_run(uname)
                r.bold = True
            else:
                r = p.add_run("")
            r.font.name = "Times New Roman"
            r.font.size = Pt(10)

    row_tot_m = t_tos_m.rows[7]
    for c_idx, cell in enumerate(row_tot_m.cells):
        cell.width = col_widths_tos[c_idx]
        set_cell_margins(cell, top=60, bottom=60, left=60, right=60)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if c_idx == 0:
            r = p.add_run("Total")
            r.bold = True
        else:
            r = p.add_run("")
            r.bold = True
        r.font.name = "Times New Roman"
        r.font.size = Pt(10)

    # Note
    p_note = doc.add_paragraph()
    p_note.paragraph_format.space_before = Pt(4)
    p_note.paragraph_format.space_after = Pt(6)
    r_n = p_note.add_run("*Note: QB approved by HOD")
    r_n.font.name = "Times New Roman"
    r_n.font.size = Pt(9.5)
    r_n.font.italic = True

    # 13. Signatures Table (Table 7)
    t_sig = doc.add_table(rows=4, cols=3)
    t_sig.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_sig.autofit = False
    set_table_borders(t_sig, color="000000", sz="4", val="single")
    sig_widths = [Inches(1.80), Inches(3.27), Inches(1.80)]
    
    sig_headers = ["", "Name of the staff / Academic Institution / Department", "Sign with date"]
    for c_idx, h in enumerate(sig_headers):
        cell = t_sig.rows[0].cells[c_idx]
        cell.width = sig_widths[c_idx]
        set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        r.font.name = "Times New Roman"
        r.font.size = Pt(10)
        r.bold = True

    roles = ["Prepared by", "Verified by", "Approved by"]
    for r_idx, role in enumerate(roles):
        row = t_sig.rows[1 + r_idx]
        for c_idx in range(3):
            row.cells[c_idx].width = sig_widths[c_idx]
            set_cell_margins(row.cells[c_idx], top=80, bottom=80, left=80, right=80)
        p0 = row.cells[0].paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r0 = p0.add_run(role)
        r0.font.name = "Times New Roman"
        r0.font.size = Pt(10)
        r0.bold = True

    # Save document
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    print(f"Successfully generated 2025 Model Examination template at: {output_path}")

if __name__ == "__main__":
    out_dir = os.path.join(os.path.dirname(__file__), "..", "templates")
    out_file = os.path.join(out_dir, "model_2025.docx")
    create_2025_model_template(out_file)
    
    # Also save in root directory
    root_file = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", "model_2025.docx"))
    create_2025_model_template(root_file)
