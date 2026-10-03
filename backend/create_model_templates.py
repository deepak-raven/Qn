import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
from docx.table import _Cell
import copy
import sys
import os

sys.path.append(os.path.dirname(os.path.abspath(__file__)))
from app.generator import sanitize_document_ooxml

def clean_cell_p(cell, text="", align=WD_ALIGN_PARAGRAPH.CENTER, bold=False, font_size=11):
    while len(cell.paragraphs) > 1:
        p_elem = cell.paragraphs[-1]._p
        p_elem.getparent().remove(p_elem)
    p = cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
    p.text = ""
    # strip numPr
    pPr = p._p.get_or_add_pPr()
    for child in list(pPr):
        if child.tag.endswith('numPr') or child.tag.endswith(':numPr') or child.tag == 'numPr':
            pPr.remove(child)
        if child.tag.endswith('pStyle') or child.tag.endswith(':pStyle') or child.tag == 'pStyle':
            pPr.remove(child)
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.08
    if text:
        run = p.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(font_size)
        run.bold = bold
    return p

def create_model_2025():
    templates_dir = os.path.join(os.path.dirname(__file__), 'templates')
    cat_2025_path = os.path.join(templates_dir, 'cat_2025.docx')
    doc = docx.Document(cat_2025_path)
    
    # 1. Header table (Table 0)
    t0 = doc.tables[0]
    tcs1 = t0.rows[1]._tr.xpath('./w:tc')
    if len(tcs1) >= 1:
        c_title = _Cell(tcs1[0], t0)
        clean_cell_p(c_title, "MODEL EXAMINATION", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=14)
        c_title.paragraphs[0].runs[0].underline = True

    # Row 3: Regulation
    tcs3 = t0.rows[3]._tr.xpath('./w:tc')
    if len(tcs3) >= 1:
        c_reg = _Cell(tcs3[0], t0)
        clean_cell_p(c_reg, "(2025-REGULATION)", align=WD_ALIGN_PARAGRAPH.CENTER, bold=False, font_size=14)

    # 2. Course Table (Table 1)
    t1 = doc.tables[1]
    clean_cell_p(t1.rows[2].cells[0], "Time: 3 Hours", align=WD_ALIGN_PARAGRAPH.LEFT, bold=True, font_size=11)
    clean_cell_p(t1.rows[2].cells[1], "Maximum Marks: 100", align=WD_ALIGN_PARAGRAPH.LEFT, bold=True, font_size=11)

    # 3. Table 3: Part A 1 Mark (10 questions)
    t3 = doc.tables[3]
    while len(t3.rows) < 11:
        t3.add_row()
    clean_cell_p(t3.rows[0].cells[0], "Q.No.", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=11)
    clean_cell_p(t3.rows[0].cells[1], "Questions", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=11)
    clean_cell_p(t3.rows[0].cells[2], "KL", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=11)
    clean_cell_p(t3.rows[0].cells[3], "CO Attainment", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=11)
    for idx in range(10):
        row = t3.rows[1 + idx]
        clean_cell_p(row.cells[0], f"{idx + 1}.", align=WD_ALIGN_PARAGRAPH.CENTER, font_size=11)
        clean_cell_p(row.cells[1], "", align=WD_ALIGN_PARAGRAPH.LEFT, font_size=11)
        clean_cell_p(row.cells[2], "K1", align=WD_ALIGN_PARAGRAPH.CENTER, font_size=11)
        default_u = f"CO{(idx // 2) + 1}"
        clean_cell_p(row.cells[3], default_u, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=11)

    # 4. Table 4: Part A Section 2 (3 Marks Short Answer - 10 questions)
    t4 = doc.tables[4]
    while len(t4.rows) < 11:
        t4.add_row()
    clean_cell_p(t4.rows[0].cells[0], "Q.No.", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=11)
    clean_cell_p(t4.rows[0].cells[1], "Questions", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=11)
    clean_cell_p(t4.rows[0].cells[2], "KL", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=11)
    clean_cell_p(t4.rows[0].cells[3], "CO Attainment", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=11)
    for idx in range(10):
        row = t4.rows[1 + idx]
        clean_cell_p(row.cells[0], f"{11 + idx}.", align=WD_ALIGN_PARAGRAPH.CENTER, font_size=11)
        clean_cell_p(row.cells[1], "", align=WD_ALIGN_PARAGRAPH.LEFT, font_size=11)
        clean_cell_p(row.cells[2], "K2", align=WD_ALIGN_PARAGRAPH.CENTER, font_size=11)
        default_u = f"CO{(idx // 2) + 1}"
        clean_cell_p(row.cells[3], default_u, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=11)

    # 5. Table 5: Part B (12 Marks Either/Or - 5 pairs = 16 rows)
    t5 = doc.tables[5]
    while len(t5.rows) < 16:
        t5.add_row()
    
    t5_hdr_tcs = t5.rows[0]._tr.xpath('./w:tc')
    if len(t5_hdr_tcs) == 4:
        clean_cell_p(_Cell(t5_hdr_tcs[0], t5), "Q.No.", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=11)
        clean_cell_p(_Cell(t5_hdr_tcs[1], t5), "Questions", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=11)
        clean_cell_p(_Cell(t5_hdr_tcs[2], t5), "KL", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=11)
        clean_cell_p(_Cell(t5_hdr_tcs[3], t5), "CO Attainment", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=11)
    elif len(t5_hdr_tcs) >= 5:
        clean_cell_p(_Cell(t5_hdr_tcs[0], t5), "Q.No.", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=11)
        clean_cell_p(_Cell(t5_hdr_tcs[1], t5), "Sub", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=11)
        clean_cell_p(_Cell(t5_hdr_tcs[2], t5), "Questions", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=11)
        clean_cell_p(_Cell(t5_hdr_tcs[3], t5), "KL", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=11)
        clean_cell_p(_Cell(t5_hdr_tcs[4], t5), "CO Attainment", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=11)

    for idx in range(5):
        row_a = t5.rows[1 + idx * 3]
        row_or = t5.rows[2 + idx * 3]
        row_b = t5.rows[3 + idx * 3]
        q_num = 21 + idx
        default_co = f"CO{idx + 1}"

        r_a_tcs = row_a._tr.xpath('./w:tc')
        if len(r_a_tcs) >= 5:
            clean_cell_p(_Cell(r_a_tcs[0], t5), f"{q_num}.", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=11)
            clean_cell_p(_Cell(r_a_tcs[1], t5), "(a)", align=WD_ALIGN_PARAGRAPH.CENTER, font_size=11)
            clean_cell_p(_Cell(r_a_tcs[2], t5), "", align=WD_ALIGN_PARAGRAPH.LEFT, font_size=11)
            clean_cell_p(_Cell(r_a_tcs[3], t5), "K3", align=WD_ALIGN_PARAGRAPH.CENTER, font_size=11)
            clean_cell_p(_Cell(r_a_tcs[4], t5), default_co, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=11)

        r_or_tcs = row_or._tr.xpath('./w:tc')
        if len(r_or_tcs) == 4:
            clean_cell_p(_Cell(r_or_tcs[0], t5), "", align=WD_ALIGN_PARAGRAPH.CENTER, font_size=11)
            clean_cell_p(_Cell(r_or_tcs[1], t5), "(OR)", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=11)
            clean_cell_p(_Cell(r_or_tcs[2], t5), "", align=WD_ALIGN_PARAGRAPH.CENTER, font_size=11)
            clean_cell_p(_Cell(r_or_tcs[3], t5), "", align=WD_ALIGN_PARAGRAPH.CENTER, font_size=11)
        elif len(r_or_tcs) >= 5:
            clean_cell_p(_Cell(r_or_tcs[0], t5), "", align=WD_ALIGN_PARAGRAPH.CENTER, font_size=11)
            clean_cell_p(_Cell(r_or_tcs[1], t5), "", align=WD_ALIGN_PARAGRAPH.CENTER, font_size=11)
            clean_cell_p(_Cell(r_or_tcs[2], t5), "(OR)", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=11)
            clean_cell_p(_Cell(r_or_tcs[3], t5), "", align=WD_ALIGN_PARAGRAPH.CENTER, font_size=11)
            clean_cell_p(_Cell(r_or_tcs[4], t5), "", align=WD_ALIGN_PARAGRAPH.CENTER, font_size=11)

        r_b_tcs = row_b._tr.xpath('./w:tc')
        if len(r_b_tcs) >= 5:
            clean_cell_p(_Cell(r_b_tcs[0], t5), "", align=WD_ALIGN_PARAGRAPH.CENTER, font_size=11)
            clean_cell_p(_Cell(r_b_tcs[1], t5), "(b)", align=WD_ALIGN_PARAGRAPH.CENTER, font_size=11)
            clean_cell_p(_Cell(r_b_tcs[2], t5), "", align=WD_ALIGN_PARAGRAPH.LEFT, font_size=11)
            clean_cell_p(_Cell(r_b_tcs[3], t5), "K3", align=WD_ALIGN_PARAGRAPH.CENTER, font_size=11)
            clean_cell_p(_Cell(r_b_tcs[4], t5), default_co, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=11)

    # 6. Table 6: TOS Questions
    t6 = doc.tables[6]
    while len(t6.rows) < 8:
        t6.add_row()
    unit_labels = ["Unit I", "Unit II", "Unit III", "Unit IV", "Unit V"]
    for u_idx in range(5):
        row = t6.rows[2 + u_idx]
        clean_cell_p(row.cells[0], unit_labels[u_idx], align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=10)
        for k_idx in range(6):
            clean_cell_p(row.cells[1 + k_idx], "", align=WD_ALIGN_PARAGRAPH.CENTER, font_size=10)
        clean_cell_p(row.cells[7], "", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=10)
    clean_cell_p(t6.rows[7].cells[0], "Total", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=10)
    for k_idx in range(6):
        clean_cell_p(t6.rows[7].cells[1 + k_idx], "", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=10)
    clean_cell_p(t6.rows[7].cells[7], "", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=10)

    # 7. Table 7: TOS Marks
    t7 = doc.tables[7]
    while len(t7.rows) < 8:
        t7.add_row()
    for u_idx in range(5):
        row = t7.rows[2 + u_idx]
        clean_cell_p(row.cells[0], unit_labels[u_idx], align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=10)
        for k_idx in range(6):
            clean_cell_p(row.cells[1 + k_idx], "", align=WD_ALIGN_PARAGRAPH.CENTER, font_size=10)
        clean_cell_p(row.cells[7], "", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=10)
    clean_cell_p(t7.rows[7].cells[0], "Total", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=10)
    for k_idx in range(6):
        clean_cell_p(t7.rows[7].cells[1 + k_idx], "", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=10)
    clean_cell_p(t7.rows[7].cells[7], "", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=10)

    # Update section headers
    for p in doc.paragraphs:
        ptxt = p.text.upper()
        if "5 X 1" in ptxt or ("PART" in ptxt and "A" in ptxt and "= 5" in ptxt):
            p.text = "PART – A (10 x 1 = 10 Marks)"
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(13)
                r.bold = True
        elif "5 X 3" in ptxt or ("PART" in ptxt and "A" in ptxt and "= 15" in ptxt):
            p.text = "PART – A (10 x 3 = 30 Marks)"
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(13)
                r.bold = True
        elif "3 X 10" in ptxt or ("PART" in ptxt and ("B" in ptxt or "C" in ptxt) and "= 30" in ptxt):
            p.text = "PART – B (5 x 12 = 60 Marks)"
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(13)
                r.bold = True

    sanitize_document_ooxml(doc)
    out_path = os.path.join(templates_dir, 'model_2025.docx')
    doc.save(out_path)
    # Also save to root
    root_path = os.path.abspath(os.path.join(templates_dir, '..', '..', 'model_2025.docx'))
    doc.save(root_path)
    print(f'Successfully generated {out_path} and {root_path}')


def create_model_2021():
    templates_dir = os.path.join(os.path.dirname(__file__), 'templates')
    cat_2025_path = os.path.join(templates_dir, 'cat_2025.docx')
    doc = docx.Document(cat_2025_path)

    # 1. Header table (Table 0)
    t0 = doc.tables[0]
    tcs1 = t0.rows[1]._tr.xpath('./w:tc')
    if len(tcs1) >= 1:
        c_title = _Cell(tcs1[0], t0)
        clean_cell_p(c_title, "MODEL EXAMINATION", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=14)
        c_title.paragraphs[0].runs[0].underline = True

    # Row 3: Regulation
    tcs3 = t0.rows[3]._tr.xpath('./w:tc')
    if len(tcs3) >= 1:
        c_reg = _Cell(tcs3[0], t0)
        clean_cell_p(c_reg, "(2021-REGULATION)", align=WD_ALIGN_PARAGRAPH.CENTER, bold=False, font_size=14)

    # 2. Course Table (Table 1)
    t1 = doc.tables[1]
    clean_cell_p(t1.rows[2].cells[0], "Time: 3 Hours", align=WD_ALIGN_PARAGRAPH.LEFT, bold=True, font_size=11)
    clean_cell_p(t1.rows[2].cells[1], "Maximum Marks: 100", align=WD_ALIGN_PARAGRAPH.LEFT, bold=True, font_size=11)

    # 3. Table 3: Part A 2 Marks (10 questions)
    t3 = doc.tables[3]
    while len(t3.rows) < 11:
        t3.add_row()
    clean_cell_p(t3.rows[0].cells[0], "Q.No.", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=11)
    clean_cell_p(t3.rows[0].cells[1], "Questions", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=11)
    clean_cell_p(t3.rows[0].cells[2], "KL", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=11)
    clean_cell_p(t3.rows[0].cells[3], "CO Attainment", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=11)
    for idx in range(10):
        row = t3.rows[1 + idx]
        clean_cell_p(row.cells[0], f"{idx + 1}.", align=WD_ALIGN_PARAGRAPH.CENTER, font_size=11)
        clean_cell_p(row.cells[1], "", align=WD_ALIGN_PARAGRAPH.LEFT, font_size=11)
        clean_cell_p(row.cells[2], "K1", align=WD_ALIGN_PARAGRAPH.CENTER, font_size=11)
        default_u = f"CO{(idx // 2) + 1}"
        clean_cell_p(row.cells[3], default_u, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=11)

    # 4. Clone 5-col Table 5 to replace Table 4 (Part B 5 pairs)
    t4 = doc.tables[4]
    t4_p = t4._tbl.getparent()
    t4_idx = t4_p.index(t4._tbl)
    t4_p.remove(t4._tbl)
    
    t_part_b_elem = copy.deepcopy(doc.tables[4]._tbl)
    t4_p.insert(t4_idx, t_part_b_elem)
    
    temp_path = os.path.join(templates_dir, 'MODEL QUESTION.docx')
    sanitize_document_ooxml(doc)
    doc.save(temp_path)
    doc = docx.Document(temp_path)
    
    # Table 4: Part B (5 pairs = 16 rows)
    t4 = doc.tables[4]
    while len(t4.rows) < 16:
        t4.add_row()
    
    t4_hdr_tcs = t4.rows[0]._tr.xpath('./w:tc')
    if len(t4_hdr_tcs) == 4:
        clean_cell_p(_Cell(t4_hdr_tcs[0], t4), "Q.No.", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=11)
        clean_cell_p(_Cell(t4_hdr_tcs[1], t4), "Questions", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=11)
        clean_cell_p(_Cell(t4_hdr_tcs[2], t4), "KL", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=11)
        clean_cell_p(_Cell(t4_hdr_tcs[3], t4), "CO Attainment", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=11)
    elif len(t4_hdr_tcs) >= 5:
        clean_cell_p(_Cell(t4_hdr_tcs[0], t4), "Q.No.", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=11)
        clean_cell_p(_Cell(t4_hdr_tcs[1], t4), "Sub", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=11)
        clean_cell_p(_Cell(t4_hdr_tcs[2], t4), "Questions", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=11)
        clean_cell_p(_Cell(t4_hdr_tcs[3], t4), "KL", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=11)
        clean_cell_p(_Cell(t4_hdr_tcs[4], t4), "CO Attainment", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=11)

    for idx in range(5):
        row_a = t4.rows[1 + idx * 3]
        row_or = t4.rows[2 + idx * 3]
        row_b = t4.rows[3 + idx * 3]
        q_num = 11 + idx
        default_co = f"CO{idx + 1}"

        r_a_tcs = row_a._tr.xpath('./w:tc')
        if len(r_a_tcs) >= 5:
            clean_cell_p(_Cell(r_a_tcs[0], t4), f"{q_num}.", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=11)
            clean_cell_p(_Cell(r_a_tcs[1], t4), "(a)", align=WD_ALIGN_PARAGRAPH.CENTER, font_size=11)
            clean_cell_p(_Cell(r_a_tcs[2], t4), "", align=WD_ALIGN_PARAGRAPH.LEFT, font_size=11)
            clean_cell_p(_Cell(r_a_tcs[3], t4), "K3", align=WD_ALIGN_PARAGRAPH.CENTER, font_size=11)
            clean_cell_p(_Cell(r_a_tcs[4], t4), default_co, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=11)

        r_or_tcs = row_or._tr.xpath('./w:tc')
        if len(r_or_tcs) == 4:
            clean_cell_p(_Cell(r_or_tcs[0], t4), "", align=WD_ALIGN_PARAGRAPH.CENTER, font_size=11)
            clean_cell_p(_Cell(r_or_tcs[1], t4), "(OR)", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=11)
            clean_cell_p(_Cell(r_or_tcs[2], t4), "", align=WD_ALIGN_PARAGRAPH.CENTER, font_size=11)
            clean_cell_p(_Cell(r_or_tcs[3], t4), "", align=WD_ALIGN_PARAGRAPH.CENTER, font_size=11)
        elif len(r_or_tcs) >= 5:
            clean_cell_p(_Cell(r_or_tcs[0], t4), "", align=WD_ALIGN_PARAGRAPH.CENTER, font_size=11)
            clean_cell_p(_Cell(r_or_tcs[1], t4), "", align=WD_ALIGN_PARAGRAPH.CENTER, font_size=11)
            clean_cell_p(_Cell(r_or_tcs[2], t4), "(OR)", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=11)
            clean_cell_p(_Cell(r_or_tcs[3], t4), "", align=WD_ALIGN_PARAGRAPH.CENTER, font_size=11)
            clean_cell_p(_Cell(r_or_tcs[4], t4), "", align=WD_ALIGN_PARAGRAPH.CENTER, font_size=11)

        r_b_tcs = row_b._tr.xpath('./w:tc')
        if len(r_b_tcs) >= 5:
            clean_cell_p(_Cell(r_b_tcs[0], t4), "", align=WD_ALIGN_PARAGRAPH.CENTER, font_size=11)
            clean_cell_p(_Cell(r_b_tcs[1], t4), "(b)", align=WD_ALIGN_PARAGRAPH.CENTER, font_size=11)
            clean_cell_p(_Cell(r_b_tcs[2], t4), "", align=WD_ALIGN_PARAGRAPH.LEFT, font_size=11)
            clean_cell_p(_Cell(r_b_tcs[3], t4), "K3", align=WD_ALIGN_PARAGRAPH.CENTER, font_size=11)
            clean_cell_p(_Cell(r_b_tcs[4], t4), default_co, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=11)

    # Table 5: Part C (1 pair = 4 rows)
    t5 = doc.tables[5]
    while len(t5.rows) > 4:
        tr = t5.rows[-1]._tr
        tr.getparent().remove(tr)
    
    t5_hdr_tcs = t5.rows[0]._tr.xpath('./w:tc')
    if len(t5_hdr_tcs) == 4:
        clean_cell_p(_Cell(t5_hdr_tcs[0], t5), "Q.No.", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=11)
        clean_cell_p(_Cell(t5_hdr_tcs[1], t5), "Questions", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=11)
        clean_cell_p(_Cell(t5_hdr_tcs[2], t5), "KL", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=11)
        clean_cell_p(_Cell(t5_hdr_tcs[3], t5), "CO Attainment", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=11)
    elif len(t5_hdr_tcs) >= 5:
        clean_cell_p(_Cell(t5_hdr_tcs[0], t5), "Q.No.", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=11)
        clean_cell_p(_Cell(t5_hdr_tcs[1], t5), "Sub", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=11)
        clean_cell_p(_Cell(t5_hdr_tcs[2], t5), "Questions", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=11)
        clean_cell_p(_Cell(t5_hdr_tcs[3], t5), "KL", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=11)
        clean_cell_p(_Cell(t5_hdr_tcs[4], t5), "CO Attainment", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=11)

    row_a = t5.rows[1]
    row_or = t5.rows[2]
    row_b = t5.rows[3]

    r5_a_tcs = row_a._tr.xpath('./w:tc')
    if len(r5_a_tcs) >= 5:
        clean_cell_p(_Cell(r5_a_tcs[0], t5), "16.", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=11)
        clean_cell_p(_Cell(r5_a_tcs[1], t5), "(a)", align=WD_ALIGN_PARAGRAPH.CENTER, font_size=11)
        clean_cell_p(_Cell(r5_a_tcs[2], t5), "", align=WD_ALIGN_PARAGRAPH.LEFT, font_size=11)
        clean_cell_p(_Cell(r5_a_tcs[3], t5), "K4", align=WD_ALIGN_PARAGRAPH.CENTER, font_size=11)
        clean_cell_p(_Cell(r5_a_tcs[4], t5), "CO5", align=WD_ALIGN_PARAGRAPH.CENTER, font_size=11)

    r5_or_tcs = row_or._tr.xpath('./w:tc')
    if len(r5_or_tcs) == 4:
        clean_cell_p(_Cell(r5_or_tcs[0], t5), "", align=WD_ALIGN_PARAGRAPH.CENTER, font_size=11)
        clean_cell_p(_Cell(r5_or_tcs[1], t5), "(OR)", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=11)
        clean_cell_p(_Cell(r5_or_tcs[2], t5), "", align=WD_ALIGN_PARAGRAPH.CENTER, font_size=11)
        clean_cell_p(_Cell(r5_or_tcs[3], t5), "", align=WD_ALIGN_PARAGRAPH.CENTER, font_size=11)
    elif len(r5_or_tcs) >= 5:
        clean_cell_p(_Cell(r5_or_tcs[0], t5), "", align=WD_ALIGN_PARAGRAPH.CENTER, font_size=11)
        clean_cell_p(_Cell(r5_or_tcs[1], t5), "", align=WD_ALIGN_PARAGRAPH.CENTER, font_size=11)
        clean_cell_p(_Cell(r5_or_tcs[2], t5), "(OR)", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=11)
        clean_cell_p(_Cell(r5_or_tcs[3], t5), "", align=WD_ALIGN_PARAGRAPH.CENTER, font_size=11)
        clean_cell_p(_Cell(r5_or_tcs[4], t5), "", align=WD_ALIGN_PARAGRAPH.CENTER, font_size=11)

    r5_b_tcs = row_b._tr.xpath('./w:tc')
    if len(r5_b_tcs) >= 5:
        clean_cell_p(_Cell(r5_b_tcs[0], t5), "", align=WD_ALIGN_PARAGRAPH.CENTER, font_size=11)
        clean_cell_p(_Cell(r5_b_tcs[1], t5), "(b)", align=WD_ALIGN_PARAGRAPH.CENTER, font_size=11)
        clean_cell_p(_Cell(r5_b_tcs[2], t5), "", align=WD_ALIGN_PARAGRAPH.LEFT, font_size=11)
        clean_cell_p(_Cell(r5_b_tcs[3], t5), "K4", align=WD_ALIGN_PARAGRAPH.CENTER, font_size=11)
        clean_cell_p(_Cell(r5_b_tcs[4], t5), "CO5", align=WD_ALIGN_PARAGRAPH.CENTER, font_size=11)

    # Table 6: TOS Questions
    t6 = doc.tables[6]
    while len(t6.rows) < 8:
        t6.add_row()
    unit_labels = ["Unit I", "Unit II", "Unit III", "Unit IV", "Unit V"]
    for u_idx in range(5):
        row = t6.rows[2 + u_idx]
        clean_cell_p(row.cells[0], unit_labels[u_idx], align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=10)
        for k_idx in range(6):
            clean_cell_p(row.cells[1 + k_idx], "", align=WD_ALIGN_PARAGRAPH.CENTER, font_size=10)
        clean_cell_p(row.cells[7], "", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=10)
    clean_cell_p(t6.rows[7].cells[0], "Total", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=10)
    for k_idx in range(6):
        clean_cell_p(t6.rows[7].cells[1 + k_idx], "", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=10)
    clean_cell_p(t6.rows[7].cells[7], "", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=10)

    # Table 7: TOS Marks
    t7 = doc.tables[7]
    while len(t7.rows) < 8:
        t7.add_row()
    for u_idx in range(5):
        row = t7.rows[2 + u_idx]
        clean_cell_p(row.cells[0], unit_labels[u_idx], align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=10)
        for k_idx in range(6):
            clean_cell_p(row.cells[1 + k_idx], "", align=WD_ALIGN_PARAGRAPH.CENTER, font_size=10)
        clean_cell_p(row.cells[7], "", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=10)
    clean_cell_p(t7.rows[7].cells[0], "Total", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=10)
    for k_idx in range(6):
        clean_cell_p(t7.rows[7].cells[1 + k_idx], "", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=10)
    clean_cell_p(t7.rows[7].cells[7], "", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=10)

    # Update section headers
    # Find all paragraphs with 'PART' in text in order
    part_paragraphs = [p for p in doc.paragraphs if "PART" in p.text.upper()]
    if len(part_paragraphs) >= 3:
        part_paragraphs[0].text = "PART – A (10 x 2 = 20 Marks)"
        part_paragraphs[1].text = "PART – B (5 x 13 = 65 Marks)"
        part_paragraphs[2].text = "PART – C (1 x 15 = 15 Marks)"
        for p in part_paragraphs[:3]:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(13)
                r.bold = True
    else:
        for p in doc.paragraphs:
            ptxt = p.text.upper()
            if "5 X 1" in ptxt or ("PART" in ptxt and "A" in ptxt and "= 5" in ptxt):
                p.text = "PART – A (10 x 2 = 20 Marks)"
            elif "5 X 3" in ptxt or ("PART" in ptxt and "A" in ptxt and "= 15" in ptxt):
                p.text = "PART – B (5 x 13 = 65 Marks)"
            elif "3 X 10" in ptxt or ("PART" in ptxt and ("B" in ptxt or "C" in ptxt) and "= 30" in ptxt):
                p.text = "PART – C (1 x 15 = 15 Marks)"
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(13)
                r.bold = True

    sanitize_document_ooxml(doc)
    out_path = os.path.join(templates_dir, 'MODEL QUESTION.docx')
    doc.save(out_path)
    # Also save to parent / root
    root_path = os.path.abspath(os.path.join(templates_dir, '..', '..', 'MODEL  QUESTION.docx'))
    doc.save(root_path)
    print(f'Successfully generated {out_path} and {root_path}')

if __name__ == '__main__':
    create_model_2025()
    create_model_2021()
