import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import os

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E1", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    tblBorders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:left w:val="none"/>
            <w:right w:val="none"/>
            <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(tblBorders)

def add_callout_box(doc, title, text, border_color="2563EB", bg_color="F8FAFC"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=160, bottom=160, left=220, right=200)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(3)
    r_title = p.add_run(f"✦ {title}\n")
    r_title.bold = True
    r_title.font.name = "Segoe UI"
    r_title.font.size = Pt(10.5)
    r_title.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
    
    r_text = p.add_run(text)
    r_text.font.name = "Segoe UI"
    r_text.font.size = Pt(9.5)
    r_text.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
    
    # Empty space after table
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(6)

def generate_doc():
    doc = docx.Document()
    
    # Page setup - 0.75 in margins
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)
    
    # --- HERO HEADER BANNER ---
    header_tbl = doc.add_table(rows=1, cols=1)
    header_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_cell = header_tbl.cell(0, 0)
    header_cell.width = Inches(7.0)
    set_cell_background(header_cell, "1E293B") # Deep Slate Navy
    set_cell_margins(header_cell, top=260, bottom=260, left=280, right=280)
    
    hp = header_cell.paragraphs[0]
    hp.paragraph_format.space_before = Pt(4)
    hp.paragraph_format.space_after = Pt(2)
    
    r_tag = hp.add_run("ANNA UNIVERSITY • REGULATION 2025 & 2021 ARCHITECTURE\n")
    r_tag.font.name = "Segoe UI"
    r_tag.font.size = Pt(8.5)
    r_tag.font.bold = True
    r_tag.font.color.rgb = RGBColor(0x38, 0xBD, 0xF8) # Cyan Accent
    
    r_head = hp.add_run("Intelligent Auto-Fill & Dual-Set Generator\n")
    r_head.font.name = "Segoe UI"
    r_head.font.size = Pt(18)
    r_head.font.bold = True
    r_head.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    
    r_sub = hp.add_run("System Overview, Regulation 2025 Blueprints, Bloom's Balancing & 0% Duplication Engine")
    r_sub.font.name = "Segoe UI"
    r_sub.font.size = Pt(10)
    r_sub.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)
    
    p_gap = doc.add_paragraph()
    p_gap.paragraph_format.space_before = Pt(4)
    p_gap.paragraph_format.space_after = Pt(8)
    
    # --- STATS BADGES (3 cards) ---
    stats_tbl = doc.add_table(rows=1, cols=3)
    stats_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    stat_data = [
        ("0.0% OVERLAP", "Strict Disjoint Cross-Set Generation", "EEF2FF", "4F46E5"),
        ("2025 REG READY", "50 Marks • 1M, 3M & 10M Layouts", "ECFDF5", "059669"),
        ("BLOOM'S BALANCED", "K1–K2 (Part A) to K4–K6 (Part C)", "FFF7ED", "D97706")
    ]
    for i, (stitle, ssub, sbg, stext_color) in enumerate(stat_data):
        c = stats_tbl.cell(0, i)
        c.width = Inches(2.25)
        set_cell_background(c, sbg)
        set_cell_margins(c, top=140, bottom=140, left=160, right=160)
        
        # Border
        tcPr = c._tc.get_or_add_tcPr()
        borders = parse_xml(f'''
            <w:tcBorders {nsdecls("w")}>
                <w:top w:val="single" w:sz="4" w:space="0" w:color="{stext_color}"/>
                <w:left w:val="none"/>
                <w:bottom w:val="none"/>
                <w:right w:val="none"/>
            </w:tcBorders>
        ''')
        tcPr.append(borders)
        
        cp = c.paragraphs[0]
        cp.paragraph_format.space_before = Pt(0)
        cp.paragraph_format.space_after = Pt(0)
        
        rt = cp.add_run(f"{stitle}\n")
        rt.bold = True
        rt.font.name = "Segoe UI"
        rt.font.size = Pt(9.5)
        hex_color = tuple(int(stext_color[j:j+2], 16) for j in (0, 2, 4))
        rt.font.color.rgb = RGBColor(*hex_color)
        
        rs = cp.add_run(ssub)
        rs.font.name = "Segoe UI"
        rs.font.size = Pt(8.5)
        rs.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
    
    p_gap2 = doc.add_paragraph()
    p_gap2.paragraph_format.space_before = Pt(6)
    p_gap2.paragraph_format.space_after = Pt(10)
    
    def add_section_title(title_text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        
        r = p.add_run(title_text)
        r.bold = True
        r.font.name = "Segoe UI"
        r.font.size = Pt(13)
        r.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A) # Dark slate
        
        # subtle divider line
        p_div = doc.add_paragraph()
        p_div.paragraph_format.space_before = Pt(0)
        p_div.paragraph_format.space_after = Pt(8)
        p_div.paragraph_format.keep_with_next = True
        r_line = p_div.add_run("―" * 48)
        r_line.font.name = "Segoe UI"
        r_line.font.size = Pt(6)
        r_line.font.color.rgb = RGBColor(0x25, 0x63, 0xEB) # Blue accent line

    # --- SECTION 1 ---
    add_section_title("1. Executive Overview & Workflow")
    
    p_intro = doc.add_paragraph()
    p_intro.paragraph_format.space_after = Pt(6)
    r = p_intro.add_run(
        "The Auto-Fill & Dual-Set Generation Engine automates the generation of compliant question papers "
        "for Anna University affiliated institutions. By analyzing uploaded question banks (.docx / .doc), "
        "the engine detects curriculum regulations (Regulation 2025 vs Regulation 2021), extracts syllabus metadata, "
        "and produces SET-I and SET-II simultaneously with mathematical zero duplication."
    )
    r.font.name = "Segoe UI"
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    add_callout_box(
        doc,
        "How Auto-Fill Works In 3 Steps",
        "1. Instant Document Ingestion: Scans uploaded QB headers for Subject Code, Subject Name, Degree, Semester, and Regulation.\n"
        "2. Stratified Sampling & Partitioning: Groups all parsed questions into Part/Unit pools and runs non-sequential Fisher-Yates randomization.\n"
        "3. Dual Disjoint Assembly: Fills SET-I using Bloom's priority rules, locks chosen IDs globally, and constructs SET-II from remaining disjoint questions.",
        border_color="2563EB",
        bg_color="F0FDF4"
    )

    # --- SECTION 2 ---
    add_section_title("2. Regulation 2025 Blueprint & Template Rules")
    
    p2025 = doc.add_paragraph()
    p2025.paragraph_format.space_after = Pt(6)
    r2025 = p2025.add_run(
        "For Regulation 2025, the exam structure has been completely redesigned into a 50-mark, 90-minute Continuous Assessment Test (CAT). "
        "The system strictly follows the 3-section layout with zero legacy Part C 15-mark questions:"
    )
    r2025.font.name = "Segoe UI"
    r2025.font.size = Pt(10)
    r2025.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    # Table for 2025
    t2025 = doc.add_table(rows=4, cols=6)
    t2025.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t2025, "CBD5E1")
    
    headers_2025 = ["Section", "Marks / Qn", "Question Count", "Pattern / Format", "Target Bloom's", "Total Marks"]
    widths_2025 = [Inches(1.1), Inches(0.9), Inches(1.1), Inches(1.8), Inches(1.1), Inches(1.0)]
    
    for j, h in enumerate(headers_2025):
        cell = t2025.cell(0, j)
        cell.width = widths_2025[j]
        set_cell_background(cell, "1E293B")
        set_cell_margins(cell, top=120, bottom=120, left=120, right=120)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j in [1, 2, 5] else WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(h)
        run.bold = True
        run.font.name = "Segoe UI"
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    rows_2025_data = [
        ("PART - A", "1 Mark", "5 (Q1–Q5)", "Direct Questions (No choice)", "K1, K2", "5 Marks"),
        ("PART - B", "3 Marks", "5 (Q6–Q10)", "Direct Short Questions", "K2, K3, K4", "15 Marks"),
        ("PART - C", "10 Marks", "3 Pairs (Q11–Q13)", "3 Either/Or Pairs (a OR b)", "K4, K5, K6", "30 Marks"),
    ]

    for i, row in enumerate(rows_2025_data):
        row_bg = "FFFFFF" if i % 2 == 0 else "F8FAFC"
        for j, val in enumerate(row):
            cell = t2025.cell(i+1, j)
            cell.width = widths_2025[j]
            set_cell_background(cell, row_bg)
            set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j in [1, 2, 5] else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(val)
            run.font.name = "Segoe UI"
            run.font.size = Pt(8.5)
            if j == 0:
                run.bold = True
                run.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
            elif j == 5:
                run.bold = True
                run.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)
            else:
                run.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

    p_gap3 = doc.add_paragraph()
    p_gap3.paragraph_format.space_before = Pt(4)
    p_gap3.paragraph_format.space_after = Pt(6)

    # --- SECTION 3 ---
    add_section_title("3. Regulation 2021 Blueprint & Dual-Unit Rule")
    
    p2021 = doc.add_paragraph()
    p2021.paragraph_format.space_after = Pt(6)
    r2021 = p2021.add_run(
        "For Regulation 2021, the system supports both Continuous Assessment Tests (50 Marks) and Model Examinations (100 Marks). "
        "It strictly enforces the Anna University 2021 Part C Dual-Unit constraint:"
    )
    r2021.font.name = "Segoe UI"
    r2021.font.size = Pt(10)
    r2021.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    add_callout_box(
        doc,
        "Anna University 2021 Part C Dual-Unit Rule",
        "In Regulation 2021 Part C (Q16 for Model Exam, Q8 for CAT), choice (a) and choice (b) MUST originate from two different curriculum units.\n"
        "• Example: Q16(a) from Unit IV and Q16(b) from Unit V.\n"
        "The engine's pick algorithms strictly inspect the unit of choice (a) and filter candidates for choice (b) from remaining distinct units.",
        border_color="D97706",
        bg_color="FFFBEB"
    )

    # --- SECTION 4 ---
    add_section_title("4. Bloom's Taxonomy Cognitive Level Balancing")

    p_bloom = doc.add_paragraph()
    p_bloom.paragraph_format.space_after = Pt(6)
    r_bl = p_bloom.add_run(
        "To ensure balanced cognitive difficulty, the auto-selection engine queries candidate pools using progressive Bloom's Taxonomy tiers:"
    )
    r_bl.font.name = "Segoe UI"
    r_bl.font.size = Pt(10)
    r_bl.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    t_bloom = doc.add_table(rows=4, cols=4)
    t_bloom.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_bloom, "CBD5E1")
    
    headers_bl = ["Section", "Target Bloom Levels", "Cognitive Actions", "Auto-Fill Selection Priority"]
    widths_bl = [Inches(1.2), Inches(1.5), Inches(2.2), Inches(2.1)]
    
    for j, h in enumerate(headers_bl):
        cell = t_bloom.cell(0, j)
        cell.width = widths_bl[j]
        set_cell_background(cell, "1E293B")
        set_cell_margins(cell, top=120, bottom=120, left=120, right=120)
        p = cell.paragraphs[0]
        run = p.add_run(h)
        run.bold = True
        run.font.name = "Segoe UI"
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    bloom_rows = [
        ("Part A (1M / 2M)", "K1 (Remember)\nK2 (Understand)", "Define, State, List, Explain, Outline, Distinguish", "Priority 1: K1/K2 fresh questions\nPriority 2: Any fresh Part A question"),
        ("Part B (3M / 13M / 16M)", "K2 (Understand)\nK3 (Apply)\nK4 (Analyze)", "Derive, Solve, Illustrate, Demonstrate, Classify, Model", "Choice (a): K2/K3/K4\nChoice (b): K3/K4/K5"),
        ("Part C (10M / 14M / 15M)", "K4 (Analyze)\nK5 (Evaluate)\nK6 (Create)", "Design, Formulate, Evaluate, Critique, Synthesize", "Choice (a): K4/K5/K6\nChoice (b): K4/K5/K6 (distinct unit)")
    ]

    for i, row in enumerate(bloom_rows):
        row_bg = "FFFFFF" if i % 2 == 0 else "F8FAFC"
        for j, val in enumerate(row):
            cell = t_bloom.cell(i+1, j)
            cell.width = widths_bl[j]
            set_cell_background(cell, row_bg)
            set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
            p = cell.paragraphs[0]
            run = p.add_run(val)
            run.font.name = "Segoe UI"
            run.font.size = Pt(8.5)
            if j == 0:
                run.bold = True
                run.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
            else:
                run.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

    p_gap4 = doc.add_paragraph()
    p_gap4.paragraph_format.space_before = Pt(4)
    p_gap4.paragraph_format.space_after = Pt(6)

    # --- SECTION 5 ---
    add_section_title("5. Document Generation & Template Switching")

    p_tmpl = doc.add_paragraph()
    p_tmpl.paragraph_format.space_after = Pt(6)
    r_tmpl = p_tmpl.add_run(
        "When generating Word documents (.docx) or displaying the interactive A4 paper preview, "
        "the system switches templates automatically based on configuration:"
    )
    r_tmpl.font.name = "Segoe UI"
    r_tmpl.font.size = Pt(10)
    r_tmpl.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    t_sw = doc.add_table(rows=4, cols=4)
    t_sw.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_sw, "CBD5E1")
    
    headers_sw = ["Regulation", "Exam Mode", "Applied Template File", "Output Layout"]
    widths_sw = [Inches(1.2), Inches(1.5), Inches(2.1), Inches(2.2)]
    
    for j, h in enumerate(headers_sw):
        cell = t_sw.cell(0, j)
        cell.width = widths_sw[j]
        set_cell_background(cell, "1E293B")
        set_cell_margins(cell, top=120, bottom=120, left=120, right=120)
        p = cell.paragraphs[0]
        run = p.add_run(h)
        run.bold = True
        run.font.name = "Segoe UI"
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    sw_data = [
        ("Regulation 2025", "CAT-1 / CAT-2 / CAT-3", "cat_2025.docx", "Part A (1M) + Part B (3M) + Part C (10M)"),
        ("Regulation 2021", "CAT-1 / CAT-2 / CAT-3", "cat_2021.docx", "Part A (2M) + Part B (13M) + Part C (14M)"),
        ("Regulation 2021", "MODEL EXAMINATION", "model_template.docx", "Part A (2M) + Part B (13/16M) + Part C (15M)")
    ]

    for i, row in enumerate(sw_data):
        row_bg = "FFFFFF" if i % 2 == 0 else "F8FAFC"
        for j, val in enumerate(row):
            cell = t_sw.cell(i+1, j)
            cell.width = widths_sw[j]
            set_cell_background(cell, row_bg)
            set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
            p = cell.paragraphs[0]
            run = p.add_run(val)
            run.font.name = "Segoe UI"
            run.font.size = Pt(8.5)
            if j == 2:
                run.bold = True
                run.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)
            else:
                run.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

    # --- FOOTER ---
    section = doc.sections[0]
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fr = fp.add_run("Automated Question Paper Generator • Official Specification Document")
    fr.font.name = "Segoe UI"
    fr.font.size = Pt(8)
    fr.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)

    output_path = r"c:\Users\karthick\Desktop\overall\questiongen\Qn\Auto_Fill_Feature_Documentation.docx"
    doc.save(output_path)
    print(f"Successfully generated DOCX at: {output_path}")

if __name__ == "__main__":
    generate_doc()
